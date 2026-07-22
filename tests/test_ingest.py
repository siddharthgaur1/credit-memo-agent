from __future__ import annotations

import pytest

from src.ingest.classify import DocType, classify
from src.ingest.parsers import ScannedPDFError, parse_file, parse_tabular


def _write_workbook(path, rows, sheet="Summary", extra_sheet=None):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    if extra_sheet:
        ws2 = wb.create_sheet(extra_sheet[0])
        for row in extra_sheet[1]:
            ws2.append(row)
    wb.save(path)
    return path


def test_excel_header_below_a_title_block_is_found(tmp_path):
    """The 'figures start on row 7 after a title block' layout."""
    path = _write_workbook(
        tmp_path / "bank_statement_summary.xlsx",
        [
            ["STATE BANK OF INDIA - ACCOUNT STATEMENT"],
            ["Meenakshi Traders"],
            ["Account number: XXXX4417  IFSC: SBIN0004512"],
            ["Statement period: Apr 2022 - Mar 2023"],
            [],
            ["Particulars", "Value"],
            ["Average monthly balance", 41.5],
            ["Total credits", 2515.0],
        ],
    )

    doc = parse_tabular(path)
    summary = doc.pages[0]

    assert summary.tables[0][0] == ["Particulars", "Value"]
    assert ["Average monthly balance", "41.5"] in summary.tables[0]
    # The title block is preserved separately, not mistaken for data.
    assert "STATE BANK OF INDIA" in summary.text


def test_multi_sheet_workbook_keeps_every_sheet_addressable(tmp_path):
    path = _write_workbook(
        tmp_path / "book.xlsx",
        [["Particulars", "Value"], ["Total credits", 100]],
        extra_sheet=("Transactions", [["Value date", "Narration"], ["01-04-2022", "NEFT inward"]]),
    )

    doc = parse_tabular(path)
    assert [p.sheet for p in doc.pages] == ["Summary", "Transactions"]
    assert all(p.number is None for p in doc.pages)


def test_merged_header_cells_are_carried_forward(tmp_path):
    path = _write_workbook(
        tmp_path / "merged.xlsx",
        [["Particulars", "FY23", None], ["Revenue", 2400, 2050]],
    )
    header = parse_tabular(path).pages[0].tables[0][0]
    assert header == ["Particulars", "FY23", "FY23"]


def test_csv_is_parsed_like_a_single_sheet(tmp_path):
    path = tmp_path / "figures.csv"
    path.write_text("Particulars,Value\nRevenue,2400\n", encoding="utf-8")

    doc = parse_file(path)
    assert doc.kind == "csv"
    assert doc.pages[0].tables[0][0] == ["Particulars", "Value"]


def test_scanned_pdf_is_reported_not_silently_empty(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "scanned_balance_sheet.pdf"
    c = canvas.Canvas(str(path))
    c.line(50, 50, 400, 700)  # ink, but no text layer and no ruled table
    c.line(400, 50, 50, 700)
    c.save()

    with pytest.raises(ScannedPDFError) as exc:
        parse_file(path)

    assert "scanned" in str(exc.value).lower()
    assert "OCR is not supported" in str(exc.value)


def test_digital_pdf_is_read_with_page_numbers(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "balance_sheet_FY23.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 760, "BALANCE SHEET AS AT 31 MARCH 2023")
    c.drawString(72, 740, "Shareholders funds")
    c.drawString(72, 720, "Total current liabilities  430.00")
    c.showPage()
    c.drawString(72, 760, "Reserves and surplus  420.00")
    c.save()

    doc = parse_file(path)
    assert [p.number for p in doc.pages] == [1, 2]
    assert "BALANCE SHEET" in doc.pages[0].text
    assert classify(doc).doc_type is DocType.BALANCE_SHEET


def test_unknown_document_asks_rather_than_guessing(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "notes.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 760, "Minutes of the site visit conducted last Tuesday afternoon.")
    c.save()

    result = classify(parse_file(path))
    assert result.doc_type is DocType.UNKNOWN
    assert result.question


def test_unsupported_file_type_is_rejected(tmp_path):
    path = tmp_path / "photo.png"
    path.write_bytes(b"not really a png")
    with pytest.raises(ValueError, match="unsupported file type"):
        parse_file(path)
